"""Update one ablation table from completed jobs; never fill pending results."""
import argparse
import json
import platform
import re
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
VARIANTS = {
    'full': 'Full SpatialESS', 'no_hash': 'Without spatial hashing',
    'coo': 'COO instead of CSR', 'no_reuse': 'Without edge-score reuse',
    'dense': 'Dense neighborhood baseline',
}


def read_time(path):
    if not path.exists():
        return {}
    text = path.read_text()
    rss = re.search(r'Maximum resident set size \(kbytes\):\s*([\d.]+)', text)
    wall = re.search(r'Elapsed \(wall clock\) time \(h:mm:ss or m:ss\):\s*(\S+)', text)
    result = {}
    if rss:
        result['peak_rss_GiB'] = float(rss[1]) / 1024**2
    if wall:
        parts = [float(p) for p in wall[1].split(':')]
        result['process_wall_seconds'] = sum(p * 60**i for i, p in enumerate(reversed(parts)))
    return result


def fmt(values):
    vals = pd.to_numeric(values, errors='coerce').dropna()
    if not len(vals):
        return 'NA'
    return f'{vals.median():.3f} [{vals.quantile(.25):.3f}, {vals.quantile(.75):.3f}]'


def main():
    (ROOT/'tables').mkdir(exist_ok=True)
    validation_rows = []
    for phase in ['tiny', 'preflight', 'formal']:
        paths = sorted((ROOT/'results'/phase).glob('*/validation.tsv'))
        if paths:
            values = pd.concat([pd.read_csv(path, sep='\t') for path in paths], ignore_index=True)
            validation_rows.append(values)
            if phase == 'preflight':
                values.to_csv(ROOT/'reports/preflight_validation.tsv', sep='\t', index=False)
    if validation_rows:
        pd.concat(validation_rows, ignore_index=True).to_csv(
            ROOT/'reports/all_validation.tsv', sep='\t', index=False)
    plan_path = ROOT/'config/run_plan.tsv'
    if not plan_path.exists():
        return
    plan = pd.read_csv(plan_path, sep='\t')
    all_runs = []
    for row in plan.to_dict('records'):
        folder = ROOT/'results/formal'/row['tag']
        summary = folder/'summary.tsv'
        if summary.exists():
            row.update(pd.read_csv(summary, sep='\t').iloc[0].to_dict())
            row.update(read_time(ROOT/'logs'/f"{row['tag']}.time.txt"))
        else:
            row['status'] = 'pending'
        validation = folder/'validation.tsv'
        if validation.exists():
            values = pd.read_csv(validation, sep='\t').iloc[0].to_dict()
            row['validation_status'] = values.pop('status')
            for key, value in values.items():
                if key not in ['case', 'phase', 'variant', 'repetition', 'nboot']:
                    row[key] = value
        else:
            row['validation_status'] = 'pending'
        exit_path = ROOT/'logs'/f"{row['tag']}.exit.json"
        if exit_path.exists():
            exit_data = json.loads(exit_path.read_text())
            row['returncode'] = exit_data['returncode']
            if exit_data['returncode'] != 0:
                row['status'] = exit_data['status']
        all_runs.append(row)
    runs = pd.DataFrame(all_runs)
    runs.to_csv(ROOT/'tables/ablation_runs.tsv', sep='\t', index=False)
    rows = []
    numeric = ['graph_seconds', 'observed_seconds', 'permutation_seconds', 'inference_seconds',
               'process_wall_seconds', 'peak_rss_GiB', 'primary_graph_array_bytes']
    for (block, case, boot, variant), g in runs.groupby(['block', 'case', 'nboot', 'variant'], sort=False):
        completed = g[g.status == 'completed'].copy()
        valid = completed[completed.validation_status == 'PASS'].copy()
        invalid_count = int((completed.validation_status == 'FAIL').sum())
        row = dict(block=block, dataset=case, N=int(g.cells.iloc[0]), permutations=int(boot),
                   configuration=VARIANTS[variant], variant=variant, completed=len(completed),
                   planned=len(g), fidelity_pass=len(valid), fidelity_fail=invalid_count)
        for col in numeric:
            values = valid[col] if col in valid else pd.Series(dtype=float)
            row[col+'_median'] = values.median() if len(values) else np.nan
            row[col+'_q25'] = values.quantile(.25) if len(values) else np.nan
            row[col+'_q75'] = values.quantile(.75) if len(values) else np.nan
            row[col+'_display'] = fmt(values)
        row['max_probability_error'] = valid.max_probability_error.max() if len(valid) else np.nan
        row['pvalues_exactly_equal'] = bool(valid.pvalues_exactly_equal.all()) if len(valid) else None
        row['failure_states'] = '; '.join(sorted(set(g.status) - {'completed', 'pending'}))
        row['status'] = ('FAIL' if invalid_count else 'COMPLETE' if len(valid) == len(g)
                         else 'INCOMPLETE' if row['failure_states'] else 'RUNNING')
        rows.append(row)
    summary = pd.DataFrame(rows)
    case_order = ['cosmx_5000', 'cosmx_10000', 'cosmx_25000',
                  'cervical_center_5000_default', 'cervical_stratified_5000_default', 'cosmx_50000']
    summary['_case_order'] = summary.dataset.map({v: i for i, v in enumerate(case_order)})
    summary['_variant_order'] = summary.variant.map({v: i for i, v in enumerate(VARIANTS)})
    summary = summary.sort_values(['_case_order', 'permutations', '_variant_order']).drop(
        columns=['_case_order', '_variant_order']).reset_index(drop=True)
    for index, row in summary.iterrows():
        refs = summary[(summary.dataset == row.dataset) &
                       (summary.permutations == row.permutations) & (summary.variant == 'full')]
        if len(refs) != 1:
            continue
        ref = refs.iloc[0]
        summary.loc[index, 'inference_time_ratio_to_full'] = (
            row.inference_seconds_median / ref.inference_seconds_median)
        summary.loc[index, 'rss_ratio_to_full'] = row.peak_rss_GiB_median / ref.peak_rss_GiB_median
        summary.loc[index, 'ratio_status'] = ('FINAL' if row.status == 'COMPLETE' and ref.status == 'COMPLETE'
                                              else 'PROVISIONAL')
    summary.to_csv(ROOT/'tables/ablation_summary.tsv', sep='\t', index=False)
    summary.to_csv(ROOT/'tables/ablation_summary.csv', index=False)
    table_rows = []
    for _, row in summary.iterrows():
        ratio = row.get('inference_time_ratio_to_full', np.nan)
        table_rows.append({
            'Dataset': row.dataset.replace('_default', ''), 'Cells': int(row.N),
            'Permutations': int(row.permutations), 'Configuration': row.configuration,
            'Graph (s)': row.graph_seconds_display, 'Observed (s)': row.observed_seconds_display,
            'Permutation (s)': row.permutation_seconds_display,
            'Inference (s)': row.inference_seconds_display,
            'Peak RSS (GiB)': row.peak_rss_GiB_display,
            'Time / Full': f'{ratio:.2f}' if pd.notna(ratio) else 'NA',
            'Validated / planned': f'{row.fidelity_pass}/{row.planned}',
            'Status': row.status + (': '+row.failure_states if row.failure_states else ''),
        })
    display = pd.DataFrame(table_rows)
    with pd.ExcelWriter(ROOT/'tables/SpatialESS_ablation_table.xlsx', engine='openpyxl') as writer:
        display.to_excel(writer, sheet_name='Ablation table', index=False)
        summary.to_excel(writer, sheet_name='Exact summaries', index=False)
        runs.to_excel(writer, sheet_name='All runs', index=False)
        notes = [
            'Times and RSS: median [25th percentile, 75th percentile] of numerical-validation PASS runs.',
            'Time / Full = ratio of median inference times, compared within the same input and nboot.',
            'Incomplete groups and ratios are provisional. Empty cells are not zero or inferred measurements.',
            'Graph, observed and permutation times exclude input loading and output serialization.',
            'Whole-process wall time and peak RSS include loading, frozen permutations and output writing.',
            'Graph payload is measured numeric-array bytes; it is not process RSS.',
            'CSR-to-COO tests graph format and traversal. CSR-to-dense tests sparse representation and traversal.',
            'No-reuse recomputes LR edge probabilities from cached invariant molecular factors and edge support.',
            'Canonical sender/receiver accumulation order, frozen label permutations and filtering are preserved.',
            'Five repetitions are technical measurements, not independent biological replicates.',
        ]
        pd.DataFrame({'Notes': notes}).to_excel(writer, sheet_name='Definitions', index=False)
        for name in ['Ablation table', 'Exact summaries', 'All runs']:
            ws = writer.sheets[name]; ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        ws = writer.sheets['Ablation table']
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(40, max(14, max(len(str(c.value)) for c in col)+2))
    doc = Document()
    section = doc.sections[0]; section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(16.54); section.page_height = Inches(11.69)
    section.left_margin = section.right_margin = Inches(.5)
    doc.styles['Normal'].font.name = 'Arial'; doc.styles['Normal'].font.size = Pt(8)
    doc.add_heading('SpatialESS component ablation', 0)
    complete = int((runs.validation_status == 'PASS').sum())
    doc.add_paragraph(f'Validated runs: {complete}/{len(runs)}. '
                      'Median [25th, 75th percentile]; only validated completed runs contribute measurements. '
                      'Incomplete rows are provisional.')
    table = doc.add_table(rows=1, cols=len(display.columns)); table.style = 'Light Shading Accent 1'
    for cell, col in zip(table.rows[0].cells, display.columns):
        cell.text = col
    for values in display.itertuples(index=False, name=None):
        for cell, value in zip(table.add_row().cells, values):
            cell.text = str(value)
    for note in notes:
        doc.add_paragraph(note)
    doc.save(ROOT/'tables/SpatialESS_ablation_table.docx')
    md = ['# SpatialESS ablation table', '',
          f'Validated: {complete}/{len(runs)} runs. Pending and incomplete values are provisional.', '',
          '| ' + ' | '.join(display.columns) + ' |',
          '| ' + ' | '.join(['---']*len(display.columns)) + ' |']
    md.extend('| '+' | '.join(str(v) for v in values)+' |'
              for values in display.itertuples(index=False, name=None))
    md.extend(['', *notes])
    (ROOT/'tables/SpatialESS_ablation_table.md').write_text('\n'.join(md)+'\n')
    result = dict(planned=len(runs), completed=int((runs.status == 'completed').sum()),
                  validated=complete, failed=int((~runs.status.isin(['completed', 'pending'])).sum()),
                  fidelity_fail=int((runs.validation_status == 'FAIL').sum()),
                  pending=int((runs.status == 'pending').sum()), table='tables/SpatialESS_ablation_table.xlsx')
    (ROOT/'reports/table_status.json').write_text(json.dumps(result, indent=2)+'\n')
    print(json.dumps(result), flush=True)


if __name__ == '__main__':
    main()
